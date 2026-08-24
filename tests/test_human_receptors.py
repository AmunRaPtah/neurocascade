"""The human PET receptor-atlas layer: Hansen 2022 ingestion (P12.1).

Every test runs offline, off the committed fixtures in
`src/neurocascade/human/fixtures/`: five byte-identical copies of the
checksum-verified parcellated tracer files plus a three-column verbatim
slice of the paper's deposited aggregation matrix. See
`fixtures/PROVENANCE.md`. A few tests additionally check the gitignored
cache (the parcel-identity xlsx, the full 19-column matrix) and are skipped
when it is absent.

The load-bearing property under test is the one the roadmap states and the
source forces: these maps are z-scored RELATIVE densities, never absolute,
and the type system refuses every read that could pretend otherwise.
"""

import os

import numpy as np
import pytest

from neurocascade.human import receptors as R
from neurocascade.human import fixtures


#: Reference values computed from the fixture bytes (which are the verified
#: published bytes) and copied here verbatim. They pin the fixture, the
#: crosswalk AND the z-scoring rule at once: if any of them moves, a test
#: fails before a downstream number can.
MOR_LACCUMB_Z = 4.198119473643431
MOR_RACCUMB_Z = 4.019926676658151
MOR_LTHAL_Z = 2.776077057490113
MOR_RTHAL_Z = 3.292653337590276
MOR_LAMYG_Z = 2.3796202644384743
MOR_LHIPPO_Z = -1.0103585391063754
D2_LPUT_Z = 19.553933928648192
D2_RPUT_Z = 20.023506959496345
D2_RCAUD_Z = 15.981719139381633
GABA_A_BZ_LACCUMB_Z = -0.6943534796226241
GABA_A_BZ_LTHAL_Z = -2.804206301697358
GABA_A_BZ_BRAINSTEM_Z = -4.587832805454584

#: The two independent carfentanil cohorts' z-scored maps agree this closely
#: over the cortex, which is what makes turtonen a replication check rather
#: than a second product map.
TURTONEN_KANTONEN_CORTEX_R = 0.9674421212076644
TURTONEN_KANTONEN_ALL83_R = 0.960458964033891


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def maps():
    return fixtures.load_fixture_receptors()


@pytest.fixture(scope="module")
def mor(maps):
    return maps["MOR"]


@pytest.fixture(scope="module")
def d2(maps):
    return maps["D2"]


@pytest.fixture(scope="module")
def gaba(maps):
    return maps["GABA_A_BZ"]


# ---------------------------------------------------------------------------
# The manifest
# ---------------------------------------------------------------------------


def test_manifest_entry_is_completely_specified():
    assert len(R.DATA_FILES) == 7
    for entry in R.DATA_FILES:
        assert len(entry.sha256) == 64, entry.key
        assert set(entry.sha256) <= set("0123456789abcdef"), entry.key
        assert entry.size_bytes > 0, entry.key
        assert entry.url.startswith("https://"), entry.key
        assert entry.filename and "/" not in entry.filename, entry.key
        assert len(entry.citation) > 40, entry.key
        assert len(entry.units) > 20, entry.key
        assert len(entry.description) > 40, entry.key
        assert entry.retrieved
        assert entry.volatility == "static"


def test_urls_are_pinned_to_a_commit():
    for entry in R.DATA_FILES:
        assert "/hansen_receptors/%s/" % R.REPOSITORY_COMMIT in entry.url
    assert R.REPOSITORY_COMMIT == "f8b41da92a733f99368c1f39d4959731aa1316c1"


def test_licence_is_recorded_on_every_entry():
    for entry in R.DATA_FILES:
        assert "CC BY-NC-SA 4.0" in entry.citation, entry.key


def test_tracer_files_cite_their_own_cohort_paper():
    for tracer in R.TRACERS:
        entry = R.data_file(tracer.manifest_key)
        assert R.TRACER_CITATIONS[tracer.key] in entry.citation, tracer.key


def test_pet_units_say_not_absolute():
    for phrase in ("NOT Bmax", "NOT fmol/mg protein", "NOT comparable"):
        assert phrase in R.RELATIVE_UNITS
    for tracer in R.TRACERS:
        units = R.data_file(tracer.manifest_key).units
        assert "never read as absolute" in units, tracer.key
    published = R.data_file(R.PUBLISHED_MATRIX_KEY)
    assert "NOT absolute" in published.units


def test_checksum_lines_are_sha256sum_c_input():
    lines = R.checksum_lines()
    assert len(lines) == len(R.DATA_FILES)
    by_filename = {entry.filename: entry.sha256 for entry in R.DATA_FILES}
    for line in lines:
        digest, filename = line.split("  ", 1)
        assert by_filename[filename] == digest


def test_missing_cache_entry_raises_instead_of_downloading(tmp_path):
    with pytest.raises(R.MissingCachedFile):
        R.ensure(R.DATA_FILES[0].key, cache_dir=str(tmp_path),
                 allow_download=False)


def test_unknown_data_file_key_names_the_alternatives():
    with pytest.raises(KeyError, match="manifest"):
        R.data_file("not_a_file")


# ---------------------------------------------------------------------------
# Parcellation and the Cammoun-033 <-> ENIGMA crosswalk
# ---------------------------------------------------------------------------


def test_parcel_counts():
    assert R.N_PARCELS == 83
    assert len(R.CORTICAL_PARCEL_INDEXES) == 68
    assert len(R.SUBCORTICAL_PARCEL_INDEXES) == 14
    assert len(set(p.label for p in R.PARCELS)) == 83
    assert R.BRAINSTEM_PARCEL.label == "brainstem"
    assert R.UNMAPPED_PARCELS == ("brainstem",)


def test_crosswalk_is_exact_for_82_parcels_and_refuses_the_rest():
    mapped = [p for p in R.PARCELS if p.connectome_label is not None]
    assert len(mapped) == 82
    assert set(R.CAMMOUN_TO_CONNECTOME.values()) - {None} == {
        p.connectome_label for p in mapped
    }
    # The brainstem's entry is an explicit None, not an absence.
    assert R.CAMMOUN_TO_CONNECTOME["brainstem"] is None
    assert R.connectome_label_for_parcel("brainstem") is None
    for cammoun_label, enigma_label in R.CAMMOUN_TO_CONNECTOME.items():
        if enigma_label is None:
            continue
        assert R.CONNECTOME_TO_CAMMOUN[enigma_label] == cammoun_label
    with pytest.raises(KeyError, match="brainstem-shaped hole|Cammoun-033"):
        R.parcel_for_connectome_label("brainstem")
    with pytest.raises(KeyError):
        R.parcel_for_connectome_label("L_nucleus_accumbens")


def test_subcortical_connectome_labels_are_the_scenarios_fourteen():
    assert R.SUBCORTICAL_CONNECTOME_LABELS == (
        "Rthal", "Rcaud", "Rput", "Rpal", "Raccumb", "Rhippo", "Ramyg",
        "Lthal", "Lcaud", "Lput", "Lpal", "Laccumb", "Lhippo", "Lamyg",
    )


def test_every_cortical_parcel_label_is_a_desikan_killiany_region():
    from neurocascade.human.atlas import DK_CORTICAL_NAMES

    for index in R.CORTICAL_PARCEL_INDEXES:
        label = R.PARCELS[index].label
        hemisphere, suffix = label.split("_", 1)
        assert hemisphere in ("L", "R"), label
        assert suffix in DK_CORTICAL_NAMES, label


# ---------------------------------------------------------------------------
# Tracer and receptor registries
# ---------------------------------------------------------------------------


def test_tracer_registry_is_unique_and_manifest_backed():
    keys = [t.key for t in R.TRACERS]
    assert len(set(keys)) == len(keys)
    for tracer in R.TRACERS:
        entry = R.data_file(tracer.manifest_key)
        assert entry.filename == tracer.manifest_key + ".csv"
        assert tracer.n_participants > 0
        assert tracer.role in ("product", "replication-check")


def test_receptor_weights_are_the_published_participant_counts():
    assert R.RECEPTORS["D2"].tracer_keys == (
        "D2_flb457_hc37_smith", "D2_flb457_hc55_sandiego",
    )
    assert R.RECEPTORS["D2"].weights == (37.0, 55.0)
    participants = [
        R._TRACERS_BY_KEY[k].n_participants for k in R.RECEPTORS["D2"].tracer_keys
    ]
    assert participants == [37, 55]


def test_not_combined_records_do_not_shadow_ingested_combinations():
    # Every recorded exclusion names a real source file and a reason.
    assert R.NOT_COMBINED, "the searched-and-not-combined record is empty"
    for key, reason in R.NOT_COMBINED.items():
        assert len(reason) > 40, key
    # turtonen is ingested but deliberately outside every product map; the
    # other recorded tracers are not ingested at all.
    ingested = {t.key for t in R.TRACERS}
    combined = set()
    for receptor in R.RECEPTORS.values():
        combined.update(receptor.tracer_keys)
    assert set(R.NOT_COMBINED) & ingested == {"MU_carfentanil_hc39_turtonen"}
    assert not (set(R.NOT_COMBINED) & combined)


def test_receptors_ingested_is_exactly_the_three_systems():
    assert R.receptors_ingested() == ("MOR", "D2", "GABA_A_BZ")


# ---------------------------------------------------------------------------
# The relative-not-absolute type system
# ---------------------------------------------------------------------------


def test_values_are_typed_never_bare_floats(mor):
    value = mor.for_connectome_label("Laccumb")
    assert isinstance(value, R.RelativeDensityValue)
    assert not isinstance(value, float)
    assert float(value) == pytest.approx(MOR_LACCUMB_Z)


def test_map_exposes_no_numeric_array_attribute(mor):
    assert not hasattr(mor, "values")
    # The array is private AND write-protected.
    with pytest.raises(AttributeError):
        setattr(mor, "_values", np.zeros(R.N_PARCELS))


def test_gain_multiplication_returns_a_plain_float(mor):
    product = 3.0 * mor.for_connectome_label("Laccumb")
    assert type(product) is float
    assert product == pytest.approx(3.0 * MOR_LACCUMB_Z)


def test_product_of_two_densities_is_refused(mor):
    a = mor.for_connectome_label("Laccumb")
    b = mor.for_connectome_label("Lthal")
    with pytest.raises(TypeError, match="product of two"):
        a * b


def test_arithmetic_with_bare_numbers_is_refused(mor):
    v = mor.for_connectome_label("Laccumb")
    with pytest.raises(TypeError):
        v + 1.0
    with pytest.raises(TypeError):
        v - 1.0
    with pytest.raises(TypeError):
        1.0 + v
    with pytest.raises(TypeError):
        v / v


def test_same_scale_arithmetic_stays_typed(mor):
    a = mor.for_connectome_label("Laccumb")
    b = mor.for_connectome_label("Lthal")
    total = a + b
    difference = a - b
    scaled = b / 2.0
    assert isinstance(total, R.RelativeDensityValue)
    assert isinstance(difference, R.RelativeDensityValue)
    assert isinstance(scaled, R.RelativeDensityValue)
    assert float(total) == pytest.approx(MOR_LACCUMB_Z + MOR_LTHAL_Z)
    assert float(-a) == pytest.approx(-MOR_LACCUMB_Z)
    assert float(abs(a)) == pytest.approx(MOR_LACCUMB_Z)


def test_cross_scale_comparison_and_arithmetic_raise(maps):
    left = maps["MOR"].for_connectome_label("Laccumb")
    # Same module scale object -> comparable. A look-alike scale is not.
    impostor = R.RelativeScale(
        name=R.NEOCORTICAL_SCALE.name,
        reference_labels=R.NEOCORTICAL_SCALE.reference_labels,
        description="field-equal look-alike",
    )
    stranger = R.RelativeDensityValue(float(left), impostor)
    assert left.scale is R.NEOCORTICAL_SCALE
    with pytest.raises(TypeError, match="different scales"):
        left < stranger
    with pytest.raises(TypeError, match="different scales"):
        left + stranger
    with pytest.raises(TypeError, match="different scales"):
        left - stranger


def test_ordering_against_numbers_is_allowed_but_hashing_is_not(mor):
    v = mor.for_connectome_label("Laccumb")
    assert v > 4.0
    assert v <= MOR_LACCUMB_Z + 1e-12
    assert v == pytest.approx(MOR_LACCUMB_Z)
    with pytest.raises(TypeError):
        hash(v)


def test_float_is_the_only_exit_from_the_type(mor):
    v = mor.for_connectome_label("Laccumb")
    assert float(v) == MOR_LACCUMB_Z
    assert "%.3f" % v == "%.3f" % MOR_LACCUMB_Z


def test_absolute_density_refusal_names_what_would_be_needed(mor):
    with pytest.raises(TypeError, match="P12.2|autoradiograph"):
        mor.absolute_density()


def test_relative_units_are_constant_and_honest(mor):
    assert mor.units is R.RELATIVE_UNITS
    assert "NOT" in mor.units
    assert mor.scale is R.NEOCORTICAL_SCALE
    assert len(mor.scale.reference_labels) == 68


def test_direct_construction_without_a_module_scale_is_refused():
    with pytest.raises(TypeError, match="RelativeScale"):
        R.PETRelativeDensityMap(
            values=np.zeros(R.N_PARCELS), receptor=R.RECEPTORS["MOR"],
            scale=None,
        )


def test_raw_parcellated_maps_have_no_arithmetic():
    a = fixtures.load_fixture_tracer("MU_carfentanil_hc204_kantonen")
    b = fixtures.load_fixture_tracer("D2_flb457_hc37_smith")
    # Different quantities on different cohort scales: the MAPS cannot be
    # combined at all -- that is what `combine_weighted` after z-scoring is
    # for. (Per-parcel reads are floats by design, with the tracer's own
    # quantity recorded on the map.)
    with pytest.raises(TypeError):
        a + b
    assert isinstance(a.for_parcel("L_accumbensarea"), float)
    # What the raw values ARE is recorded, not assumed.
    assert a.tracer.quantity == "BPnd"
    assert "Vt-derived" in b.tracer.quantity or b.tracer.quantity == "BPnd"


# ---------------------------------------------------------------------------
# Loaders, off the fixtures
# ---------------------------------------------------------------------------


def test_loaders_produce_all_three_systems_on_one_scale(maps):
    assert sorted(maps) == ["D2", "GABA_A_BZ", "MOR"]
    for key, m in maps.items():
        assert isinstance(m, R.PETRelativeDensityMap), key
        assert m.scale is R.NEOCORTICAL_SCALE, key
        assert np.all(np.isfinite([float(v) for _, v in m.measured()]))


def test_mor_resolves_over_the_subcortical_scenarios_parcels(mor):
    values = mor.subcortical_values()
    assert set(values) == set(R.SUBCORTICAL_CONNECTOME_LABELS)
    assert values["Laccumb"] == pytest.approx(MOR_LACCUMB_Z)
    assert values["Raccumb"] == pytest.approx(MOR_RACCUMB_Z)
    assert values["Lthal"] == pytest.approx(MOR_LTHAL_Z)
    assert values["Rthal"] == pytest.approx(MOR_RTHAL_Z)
    assert values["Lamyg"] == pytest.approx(MOR_LAMYG_Z)
    assert values["Lhippo"] == pytest.approx(MOR_LHIPPO_Z)


def test_d2_and_gaba_resolve_over_the_same_parcels(d2, gaba):
    d2_values = d2.subcortical_values()
    gaba_values = gaba.subcortical_values()
    assert d2_values["Lput"] == pytest.approx(D2_LPUT_Z)
    assert d2_values["Rput"] == pytest.approx(D2_RPUT_Z)
    assert d2_values["Rcaud"] == pytest.approx(D2_RCAUD_Z)
    assert gaba_values["Laccumb"] == pytest.approx(GABA_A_BZ_LACCUMB_Z)
    assert gaba_values["Lthal"] == pytest.approx(GABA_A_BZ_LTHAL_Z)
    # brainstem has no ENIGMA label but IS measured; it stays reachable.
    assert float(gaba.for_parcel("brainstem")) == pytest.approx(
        GABA_A_BZ_BRAINSTEM_Z
    )


def test_published_d2_column_is_reproduced_bit_for_bit(d2):
    published = fixtures.load_fixture_published_column("D2")
    ours = np.array(
        [float(d2.for_parcel(R.PARCELS[i].label)) for i in R._PUBLISHED_ROWS]
    )
    assert ours.shape == (68,)
    assert np.all(ours == published)


def test_published_single_tracer_columns_match_the_raw_maps_bit_for_bit():
    published_mor = fixtures.load_fixture_published_column("MOR")
    published_gaba = fixtures.load_fixture_published_column("GABA_A_BZ")
    raw_mor = fixtures.load_fixture_tracer("MU_carfentanil_hc204_kantonen")
    raw_gaba = fixtures.load_fixture_tracer("GABAa-bz_flumazenil_hc16_norgaard")
    ours_mor = np.array(
        [raw_mor.for_parcel(R.PARCELS[i].label) for i in R._PUBLISHED_ROWS]
    )
    ours_gaba = np.array(
        [raw_gaba.for_parcel(R.PARCELS[i].label) for i in R._PUBLISHED_ROWS]
    )
    assert np.all(ours_mor == published_mor)
    assert np.all(ours_gaba == published_gaba)


def test_deposit_scale_registry_matches_the_verified_heterogeneity():
    assert set(R.PUBLISHED_COLUMN_SCALES) == {"D2", "MOR", "GABA_A_BZ"}
    assert "z-scored" in R.PUBLISHED_COLUMN_SCALES["D2"]
    for single in ("MOR", "GABA_A_BZ"):
        assert "raw" in R.PUBLISHED_COLUMN_SCALES[single]
        assert "not z-scored" in R.PUBLISHED_COLUMN_SCALES[single]


def test_d2_weighted_average_algebra(d2):
    """The combination is exactly (z_smith*37 + z_sandiego*55)/92."""
    smith = fixtures.load_fixture_tracer("D2_flb457_hc37_smith")
    sandiego = fixtures.load_fixture_tracer("D2_flb457_hc55_sandiego")

    def z(raw, label):
        m, sd = raw.reference_population_mean_and_sd()
        return (raw.for_parcel(label) - m) / sd

    for label in ("L_putamen", "R_caudate", "L_accumbensarea", "brainstem"):
        expected = (37.0 * z(smith, label) + 55.0 * z(sandiego, label)) / 92.0
        got = d2.for_parcel(label)
        assert float(got) == float(expected), label


def test_combine_weighted_records_the_whole_bill_of_materials(d2):
    assert d2.receptor.tracer_keys == (
        "D2_flb457_hc37_smith", "D2_flb457_hc55_sandiego",
    )
    assert d2.receptor.weights == (37.0, 55.0)
    assert "published" in d2.receptor.combination.lower()
    # A single-tracer map says so, and does not claim a combination.
    solo = fixtures.load_fixture_receptor("MOR")
    assert solo.receptor.tracer_keys == ("MU_carfentanil_hc204_kantonen",)


def test_combine_weighted_refuses_bad_inputs(maps):
    smith = fixtures.load_fixture_tracer("D2_flb457_hc37_smith").relative_to_cortex()
    sandiego = (
        fixtures.load_fixture_tracer("D2_flb457_hc55_sandiego")
        .relative_to_cortex()
    )
    with pytest.raises(ValueError, match="weights"):
        R.combine_weighted([smith], [])
    with pytest.raises(ValueError, match="non-negative"):
        R.combine_weighted([smith, sandiego], [37.0, -1.0])
    with pytest.raises(ValueError, match="different receptors"):
        R.combine_weighted([smith, maps["MOR"]], [1.0, 1.0])
    impostor = R.RelativeScale(
        name=smith.scale.name,
        reference_labels=smith.scale.reference_labels,
        description="look-alike",
    )
    stranger = R.PETRelativeDensityMap(
        values=[float(v) for _, v in smith.measured()],
        receptor=R.RECEPTORS["D2"],
        scale=impostor,
    )
    with pytest.raises(TypeError, match="different scales"):
        R.combine_weighted([smith, stranger], [1.0, 1.0])


# ---------------------------------------------------------------------------
# Anatomy: the maps must say what the literature says they say
# ---------------------------------------------------------------------------


def test_anatomy_mor_peaks_in_accumbens_and_thalamus(mor):
    top = [label for label, _ in mor.top(6)]
    assert "L_accumbensarea" == top[0]
    assert "R_accumbensarea" == top[1]
    assert any("thalamusproper" in label for label in top)
    accumbens = max(
        float(mor.for_connectome_label(lab))
        for lab in ("Laccumb", "Raccumb")
    )
    cortex = max(float(mor.for_parcel(R.PARCELS[i].label))
                 for i in R.CORTICAL_PARCEL_INDEXES)
    assert accumbens > cortex


def test_anatomy_d2_is_striatum_dominant(d2):
    top = [label for label, _ in d2.top(6)]
    for label in top:
        assert any(
            nucleus in label
            for nucleus in ("putamen", "caudate", "pallidum", "accumbens")
        ), top
    striatum = min(
        float(d2.for_connectome_label(lab))
        for lab in ("Lput", "Rput", "Lcaud", "Rcaud")
    )
    cortex_max = max(float(d2.for_parcel(R.PARCELS[i].label))
                     for i in R.CORTICAL_PARCEL_INDEXES)
    assert striatum > cortex_max > 2.0


def test_anatomy_gaba_bz_is_cortical_with_a_dead_brainstem(gaba):
    brainstem = float(gaba.for_parcel("brainstem"))
    subcortex = [float(v) for v in gaba.subcortical_values().values()]
    cortex_max = max(float(gaba.for_parcel(R.PARCELS[i].label))
                     for i in R.CORTICAL_PARCEL_INDEXES)
    assert brainstem == pytest.approx(GABA_A_BZ_BRAINSTEM_Z)
    assert brainstem < min(subcortex) < 0.0 < cortex_max
    # Benzodiazepine binding peaks in occipital/visual cortex, as published.
    assert "pericalcarine" in gaba.top(1)[0][0]


def test_the_deposited_zero_at_r_pericalcarine_is_kept():
    """A real feature of the kantonen deposit, not cleaned up.

    The aggregation matrix carries the same zero (bit-for-bit test above),
    so the publication treats it as a measurement and so does this module.
    """
    raw = fixtures.load_fixture_tracer("MU_carfentanil_hc204_kantonen")
    assert raw.for_parcel("R_pericalcarine") == 0.0


# ---------------------------------------------------------------------------
# Replication check: the second carfentanil cohort
# ---------------------------------------------------------------------------


def test_turtonen_replicates_kantonen_but_is_not_a_product_map():
    kan = fixtures.load_fixture_tracer("MU_carfentanil_hc204_kantonen")
    tur = fixtures.load_fixture_tracer("MU_carfentanil_hc39_turtonen")
    kv = np.array([kan.for_parcel(p.label) for p in R.PARCELS])
    tv = np.array([tur.for_parcel(p.label) for p in R.PARCELS])
    cortical = np.array(R.CORTICAL_PARCEL_INDEXES)
    r_cortex = float(np.corrcoef(kv[cortical], tv[cortical])[0, 1])
    r_all = float(np.corrcoef(kv, tv)[0, 1])
    assert r_cortex == pytest.approx(TURTONEN_KANTONEN_CORTEX_R)
    assert r_all == pytest.approx(TURTONEN_KANTONEN_ALL83_R)
    assert r_cortex > 0.95
    # And it stays out of the product map: provenance says replication.
    assert R._TRACERS_BY_KEY["MU_carfentanil_hc39_turtonen"].role == \
        "replication-check"
    solo = fixtures.load_fixture_tracer("MU_carfentanil_hc39_turtonen")
    assert "replication" in solo.relative_to_cortex().receptor.combination


# ---------------------------------------------------------------------------
# The connectome join
# ---------------------------------------------------------------------------


def test_coupling_over_connectome_covers_every_mapped_region(maps):
    labels = sorted(R.CONNECTOME_TO_CAMMOUN)
    coupling = R.coupling_over_connectome(maps["MOR"], mechanism_gains=-3.0,
                                          labels=labels)
    assert sorted(coupling) == labels
    assert coupling["Laccumb"] == pytest.approx(-3.0 * MOR_LACCUMB_Z)
    for label in labels:
        expected = -3.0 * float(maps["MOR"].for_connectome_label(label))
        assert coupling[label] == pytest.approx(expected), label
    # products are plain floats: coupling coefficients, not densities
    assert all(type(v) is float for v in coupling.values())


def test_coupling_over_connectome_accepts_the_multi_system_shape(maps):
    labels = list(R.SUBCORTICAL_CONNECTOME_LABELS)
    nodes = {label: label for label in labels}
    coupling = R.coupling_over_connectome(maps, mechanism_gains=-2.0,
                                          labels=labels)
    assert sorted(coupling) == ["D2", "GABA_A_BZ", "MOR"]
    for key, per_node in coupling.items():
        assert sorted(per_node) == sorted(labels), key
        for label in labels:
            expected = -2.0 * float(maps[key].for_connectome_label(label))
            assert per_node[label] == pytest.approx(expected)


def test_coupling_refuses_an_uncovered_region(maps):
    with pytest.raises(KeyError, match="Cammoun-033"):
        R.coupling_over_connectome(maps["MOR"], mechanism_gains=1.0,
                                   labels=["brainstem"])


# ---------------------------------------------------------------------------
# Parse honesty
# ---------------------------------------------------------------------------


def test_wrong_length_file_is_rejected(tmp_path):
    path = tmp_path / "short.csv"
    path.write_text(",".join(["1.0"] * 82))
    with pytest.raises(ValueError, match="83"):
        R.read_parcellated(str(path), "MU_carfentanil_hc204_kantonen")


def test_non_finite_cell_is_rejected_not_zero_filled(tmp_path):
    values = ["1.0"] * 83
    values[10] = "NaN"
    path = tmp_path / "nan.csv"
    path.write_text(",".join(values))
    with pytest.raises(ValueError, match="non-finite"):
        R.read_parcellated(str(path), "MU_carfentanil_hc204_kantonen")


def test_unknown_tracer_and_unknown_parcel_name_the_knowns(tmp_path):
    path = tmp_path / "row.csv"
    path.write_text(",".join(["1.0"] * 83))
    with pytest.raises(KeyError, match="MU_carfentanil"):
        R.read_parcellated(str(path), "MU_carfentanil_hc999_nobody")
    raw = fixtures.load_fixture_tracer("MU_carfentanil_hc204_kantonen")
    with pytest.raises(KeyError, match="Cammoun-033 parcel labels"):
        raw.for_parcel("Laccumb")


def test_published_column_slice_reader_validates_shape(tmp_path):
    good = fixtures.FIXTURE_PUBLISHED_MATRIX_PATH
    column = R.read_published_column_slice(good, "D2")
    assert column.shape == (68,)
    with pytest.raises(KeyError, match="fixture slice carries"):
        R.read_published_column_slice(good, "H1")
    bad = tmp_path / "bad.csv"
    bad.write_text(",".join(["1.0", "2.0"]))
    with pytest.raises(ValueError, match="expected shape"):
        R.read_published_column_slice(str(bad), "MOR")


# ---------------------------------------------------------------------------
# The full cached files, when present (optional, skipped offline)
# ---------------------------------------------------------------------------


def test_cached_files_verify_against_the_manifest():
    try:
        paths = [R.cache_path(entry.key) for entry in R.DATA_FILES]
    except Exception:
        pytest.skip("cache path unavailable")
    if not all(os.path.exists(p) for p in paths):
        pytest.skip("data/human/ cache absent; run scripts/fetch_human_receptors.py")
    for entry in R.DATA_FILES:
        R.verify(entry.key)


def test_hard_coded_parcels_table_matches_the_published_xlsx():
    pytest.importorskip("openpyxl")
    path = R.cache_path("hansen_lausanne_region_info_scale033")
    if not os.path.exists(path):
        pytest.skip(
            "data/human/ cache absent; run scripts/fetch_human_receptors.py"
        )
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(cell).strip().lower() if cell is not None else ""
              for cell in rows[0]]
    id_col = header.index("id")
    label_col = header.index("label")
    hemisphere_col = header.index("hemisphere")
    structure_col = header.index("structure")
    entries = rows[1:]
    assert len(entries) == R.N_PARCELS
    for row, parcel in zip(entries, R.PARCELS):
        name = str(row[label_col]).strip()
        hemisphere = str(row[hemisphere_col]).strip()
        # The published table stores the bare region name plus a hemisphere
        # column; the parcellated CSVs' labels are their composition. The
        # brainstem row is the one exception: it is a single midline
        # structure whose 'L' is a placeholder, and this module keeps the
        # bare name rather than inventing a hemisphere.
        expected_label = (
            "brainstem" if name == "brainstem" else "%s_%s" % (hemisphere, name)
        )
        assert int(row[id_col]) == parcel.id, parcel.label
        assert expected_label == parcel.label, parcel.label
        assert hemisphere == parcel.hemisphere, parcel.label
        assert str(row[structure_col]).strip() == parcel.structure, parcel.label


def test_cached_loaders_agree_with_the_fixture_builds():
    fixture_maps = fixtures.load_fixture_receptors()
    try:
        cached = R.load_receptors(cache_dir=R.DEFAULT_CACHE_DIR)
    except R.MissingCachedFile:
        pytest.skip("data/human/ cache absent; run scripts/fetch_human_receptors.py")
    for key in R.receptors_ingested():
        for parcel in R.PARCELS:
            assert float(cached[key].for_parcel(parcel.label)) == float(
                fixture_maps[key].for_parcel(parcel.label)
            ), (key, parcel.label)


def test_full_cached_matrix_columns_match_the_fixture_slice():
    full_path = R.cache_path(R.PUBLISHED_MATRIX_KEY)
    slice_path = fixtures.FIXTURE_PUBLISHED_MATRIX_PATH
    if not os.path.exists(full_path):
        pytest.skip("data/human/ cache absent; run scripts/fetch_human_receptors.py")
    for key in R.PUBLISHED_FIXTURE_COLUMNS:
        full = R.read_published_cortical_columns(full_path, key)
        sliced = R.read_published_column_slice(slice_path, key)
        assert np.all(full == sliced), key
